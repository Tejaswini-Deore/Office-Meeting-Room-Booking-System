import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "meeting_room_bookings.xlsx"


# ---------------- CREATE EXCEL FILE ----------------

def create_excel_file():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Bookings"

        headers = [
            "Booking ID",
            "Employee Name",
            "Department",
            "Room Name",
            "Meeting Date",
            "Start Time",
            "End Time",
            "Number of People",
            "Purpose"
        ]

        ws.append(headers)
        wb.save(FILE_NAME)


# ---------------- LOGIN WINDOW ----------------

def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":
        login_window.destroy()
        open_dashboard()
    else:
        messagebox.showerror("Login Error", "Invalid username or password")


# ---------------- DASHBOARD ----------------

def open_dashboard():

    dashboard = tk.Tk()
    dashboard.title("Office Meeting Room Booking System")
    dashboard.geometry("900x600")
    dashboard.config(bg="#f2f2f2")

    title = tk.Label(
        dashboard,
        text="Office Meeting Room Booking System",
        font=("Arial", 22, "bold"),
        bg="#f2f2f2"
    )
    title.pack(pady=20)

    button_frame = tk.Frame(dashboard, bg="#f2f2f2")
    button_frame.pack(pady=20)

    tk.Button(
        button_frame,
        text="Book Meeting Room",
        width=22,
        height=2,
        command=book_room
    ).grid(row=0, column=0, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="View Bookings",
        width=22,
        height=2,
        command=view_bookings
    ).grid(row=0, column=1, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Update Booking",
        width=22,
        height=2,
        command=update_booking
    ).grid(row=1, column=0, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Cancel Booking",
        width=22,
        height=2,
        command=delete_booking
    ).grid(row=1, column=1, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Search Booking",
        width=22,
        height=2,
        command=search_booking
    ).grid(row=2, column=0, padx=10, pady=10)

    tk.Button(
        button_frame,
        text="Exit",
        width=22,
        height=2,
        command=dashboard.destroy
    ).grid(row=2, column=1, padx=10, pady=10)

    dashboard.mainloop()


# ---------------- BOOK ROOM ----------------

def book_room():

    window = tk.Toplevel()
    window.title("Book Meeting Room")
    window.geometry("500x600")

    tk.Label(
        window,
        text="Book Meeting Room",
        font=("Arial", 18, "bold")
    ).pack(pady=15)

    fields = [
        "Booking ID",
        "Employee Name",
        "Department",
        "Room Name",
        "Meeting Date",
        "Start Time",
        "End Time",
        "Number of People",
        "Purpose"
    ]

    entries = {}

    for field in fields:

        frame = tk.Frame(window)
        frame.pack(pady=5)

        tk.Label(
            frame,
            text=field + ":",
            width=20,
            anchor="w"
        ).pack(side="left")

        entry = tk.Entry(frame, width=30)
        entry.pack(side="left")

        entries[field] = entry

    def save_booking():

        values = [entries[field].get() for field in fields]

        if "" in values:
            messagebox.showwarning(
                "Warning",
                "Please fill all fields."
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb["Bookings"]

        # Check duplicate Booking ID
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == values[0]:
                messagebox.showerror(
                    "Error",
                    "Booking ID already exists."
                )
                return

        ws.append(values)
        wb.save(FILE_NAME)

        messagebox.showinfo(
            "Success",
            "Meeting room booked successfully!"
        )

        window.destroy()

    tk.Button(
        window,
        text="Save Booking",
        width=20,
        command=save_booking
    ).pack(pady=20)


# ---------------- VIEW BOOKINGS ----------------

def view_bookings():

    window = tk.Toplevel()
    window.title("All Meeting Room Bookings")
    window.geometry("1100x500")

    columns = (
        "Booking ID",
        "Employee Name",
        "Department",
        "Room Name",
        "Meeting Date",
        "Start Time",
        "End Time",
        "People",
        "Purpose"
    )

    tree = ttk.Treeview(
        window,
        columns=columns,
        show="headings"
    )

    for column in columns:
        tree.heading(column, text=column)
        tree.column(column, width=110)

    tree.pack(fill="both", expand=True)

    wb = load_workbook(FILE_NAME)
    ws = wb["Bookings"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", "end", values=row)


# ---------------- UPDATE BOOKING ----------------

def update_booking():

    window = tk.Toplevel()
    window.title("Update Booking")
    window.geometry("500x650")

    tk.Label(
        window,
        text="Update Meeting Booking",
        font=("Arial", 18, "bold")
    ).pack(pady=15)

    tk.Label(window, text="Enter Booking ID").pack()

    id_entry = tk.Entry(window)
    id_entry.pack(pady=5)

    fields = [
        "Employee Name",
        "Department",
        "Room Name",
        "Meeting Date",
        "Start Time",
        "End Time",
        "Number of People",
        "Purpose"
    ]

    entries = {}

    for field in fields:

        frame = tk.Frame(window)
        frame.pack(pady=5)

        tk.Label(
            frame,
            text=field + ":",
            width=20,
            anchor="w"
        ).pack(side="left")

        entry = tk.Entry(frame, width=30)
        entry.pack(side="left")

        entries[field] = entry

    def update_data():

        booking_id = id_entry.get()

        if booking_id == "":
            messagebox.showwarning(
                "Warning",
                "Enter Booking ID."
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb["Bookings"]

        found = False

        for row in ws.iter_rows(min_row=2):

            if str(row[0].value) == booking_id:

                found = True

                values = [
                    entries[field].get()
                    for field in fields
                ]

                for i, value in enumerate(values, start=2):
                    if value != "":
                        row[i - 1].value = value

                wb.save(FILE_NAME)

                messagebox.showinfo(
                    "Success",
                    "Booking updated successfully!"
                )

                window.destroy()
                break

        if not found:
            messagebox.showerror(
                "Error",
                "Booking ID not found."
            )

    tk.Button(
        window,
        text="Update Booking",
        width=20,
        command=update_data
    ).pack(pady=20)


# ---------------- DELETE BOOKING ----------------

def delete_booking():

    window = tk.Toplevel()
    window.title("Cancel Booking")
    window.geometry("400x250")

    tk.Label(
        window,
        text="Cancel Meeting Booking",
        font=("Arial", 18, "bold")
    ).pack(pady=20)

    tk.Label(
        window,
        text="Enter Booking ID:"
    ).pack()

    id_entry = tk.Entry(window)
    id_entry.pack(pady=10)

    def delete_data():

        booking_id = id_entry.get()

        wb = load_workbook(FILE_NAME)
        ws = wb["Bookings"]

        found = False

        for row in range(2, ws.max_row + 1):

            if str(ws.cell(row, 1).value) == booking_id:

                ws.delete_rows(row, 1)
                wb.save(FILE_NAME)

                found = True

                messagebox.showinfo(
                    "Success",
                    "Booking cancelled successfully!"
                )

                window.destroy()
                break

        if not found:
            messagebox.showerror(
                "Error",
                "Booking ID not found."
            )

    tk.Button(
        window,
        text="Cancel Booking",
        width=20,
        command=delete_data
    ).pack(pady=15)


# ---------------- SEARCH BOOKING ----------------

def search_booking():

    window = tk.Toplevel()
    window.title("Search Booking")
    window.geometry("1100x500")

    tk.Label(
        window,
        text="Search Booking by Employee Name",
        font=("Arial", 16, "bold")
    ).pack(pady=10)

    search_entry = tk.Entry(window, width=40)
    search_entry.pack(pady=5)

    columns = (
        "Booking ID",
        "Employee Name",
        "Department",
        "Room Name",
        "Meeting Date",
        "Start Time",
        "End Time",
        "People",
        "Purpose"
    )

    tree = ttk.Treeview(
        window,
        columns=columns,
        show="headings"
    )

    for column in columns:
        tree.heading(column, text=column)
        tree.column(column, width=110)

    tree.pack(fill="both", expand=True, pady=10)

    def search():

        for item in tree.get_children():
            tree.delete(item)

        name = search_entry.get().lower()

        wb = load_workbook(FILE_NAME)
        ws = wb["Bookings"]

        for row in ws.iter_rows(min_row=2, values_only=True):

            if name in str(row[1]).lower():
                tree.insert(
                    "",
                    "end",
                    values=row
                )

    tk.Button(
        window,
        text="Search",
        width=15,
        command=search
    ).pack(pady=5)


# ---------------- MAIN PROGRAM ----------------

create_excel_file()

login_window = tk.Tk()
login_window.title("Office Meeting Room Booking System - Login")
login_window.geometry("450x350")

tk.Label(
    login_window,
    text="Office Meeting Room Booking System",
    font=("Arial", 18, "bold")
).pack(pady=30)

tk.Label(
    login_window,
    text="Username"
).pack()

username_entry = tk.Entry(login_window, width=30)
username_entry.pack(pady=5)

tk.Label(
    login_window,
    text="Password"
).pack()

password_entry = tk.Entry(
    login_window,
    width=30,
    show="*"
)
password_entry.pack(pady=5)

tk.Button(
    login_window,
    text="Login",
    width=20,
    command=login
).pack(pady=20)

tk.Label(
    login_window,
    text="Username: admin   Password: 1234",
    fg="gray"
).pack()

login_window.mainloop()
